"""Build a Word document from contract text (LLM markdown output)."""

from __future__ import annotations

import re
from io import BytesIO

from docx import Document
from docx.shared import Pt

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^(\s*)([-*+])\s+(.*)$")
_NUMBERED_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
_HRULE_RE = re.compile(r"^(-{3,}|\*{3,}|_{3,})\s*$")


def _strip_codeblock_fence(text: str) -> str:
    t = (text or "").strip()
    if not t.startswith("```"):
        return t
    lines = t.split("\n")
    if lines and lines[0].startswith("```"):
        lines = lines[1:]
    if lines and lines[-1].strip() == "```":
        lines = lines[:-1]
    return "\n".join(lines).strip()


def _heading_level(marker: str) -> int:
    return min(len(marker), 9)


def _add_inline_runs(paragraph, text: str) -> None:
    """Parse **bold**, *italic*, and `code` into Word runs."""
    if not text:
        return
    pattern = re.compile(
        r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*]+\*|_[^_]+_|`[^`]+`)"
    )
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("__") and token.endswith("__"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("_") and token.endswith("_"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(token)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _add_paragraph(doc: Document, text: str, *, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    _add_inline_runs(paragraph, text.strip())


def _add_heading(doc: Document, level: int, text: str) -> None:
    clean = text.strip()
    if not clean:
        doc.add_paragraph("")
        return
    doc.add_heading(clean, level=_heading_level("#" * level))


def contract_text_to_docx_bytes(body: str) -> bytes:
    """Turn markdown contract text into a .docx with headings, lists, and emphasis."""
    doc = Document()
    normalized = _strip_codeblock_fence(body or "")
    normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")
    lines = normalized.split("\n")

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if _HRULE_RE.match(stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("─" * 40)
            run.font.size = Pt(8)
            i += 1
            continue

        heading = _HEADING_RE.match(stripped)
        if heading:
            _add_heading(doc, len(heading.group(1)), heading.group(2))
            i += 1
            continue

        bullet = _BULLET_RE.match(line)
        if bullet:
            _add_paragraph(doc, bullet.group(3), style="List Bullet")
            i += 1
            continue

        numbered = _NUMBERED_RE.match(line)
        if numbered:
            _add_paragraph(doc, numbered.group(3), style="List Number")
            i += 1
            continue

        # Collect consecutive plain lines into one paragraph.
        block: list[str] = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            nxt_stripped = nxt.strip()
            if (
                not nxt_stripped
                or _HEADING_RE.match(nxt_stripped)
                or _BULLET_RE.match(nxt)
                or _NUMBERED_RE.match(nxt)
                or _HRULE_RE.match(nxt_stripped)
            ):
                break
            block.append(nxt_stripped)
            i += 1
        _add_paragraph(doc, " ".join(block))

    if not doc.paragraphs:
        doc.add_paragraph("")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
