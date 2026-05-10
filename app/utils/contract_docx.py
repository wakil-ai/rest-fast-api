"""Build a Word document from plain contract text (LLM output)."""

from __future__ import annotations

from io import BytesIO

from docx import Document


def contract_text_to_docx_bytes(body: str) -> bytes:
    """Turn plain text (optionally multi-line) into a minimal .docx file."""
    doc = Document()
    normalized = (body or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = normalized.split("\n")
    if not lines:
        doc.add_paragraph("")
    else:
        for line in lines:
            doc.add_paragraph(line)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
